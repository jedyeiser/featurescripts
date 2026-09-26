"""What changed in a deck since Claude generated it: slide by slide text, images, layout and speaker notes.

The generated deck is kept as docs/decks/<slug>/.baseline/<slug>.generated.pptx (copied at every build). The user
edits docs/decks/<slug>/<slug>.pptx in PowerPoint; this reports the differences so they can be folded back into
the spec / the deck builder (and so later decks inherit style changes).

Instructions to Claude can go in a slide's speaker notes (any text), or in a PowerPoint comment -- both reported.

usage (repo root):  python docs/tooling/deck_diff.py split_plus
"""
import hashlib
import sys
import zipfile
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent.parent


def inches(v):
    return round(Emu(v).inches, 2) if v is not None else None


def describe(prs):
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        shapes = []
        for sh in slide.shapes:
            entry = {"kind": sh.shape_type and str(sh.shape_type).split(".")[-1].split(" ")[0] or "shape",
                     "box": (inches(sh.left), inches(sh.top), inches(sh.width), inches(sh.height))}
            if sh.has_text_frame and sh.text_frame.text.strip():
                entry["text"] = sh.text_frame.text.strip()
            if getattr(sh, "has_table", False) and sh.has_table:
                entry["table"] = [[c.text for c in r.cells] for r in sh.table.rows]
            if sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
                entry["image"] = hashlib.sha1(sh.image.blob).hexdigest()[:10]
            shapes.append(entry)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        title = next((s["text"].splitlines()[0] for s in shapes if "text" in s), "(no text)")
        slides.append({"n": idx, "title": title, "shapes": shapes, "notes": notes})
    return slides


def comments(path):
    """(slide number, text) for every PowerPoint comment, mapped through each slide's relationships."""
    prs = Presentation(path)
    order = {sl.part.partname.split("/")[-1]: i + 1 for i, sl in enumerate(prs.slides)}
    out = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            m = re.match(r"ppt/slides/_rels/(slide\d+\.xml)\.rels$", name)
            if not m:
                continue
            for target in re.findall(r'Target="\.\./comments/([^"]+)"', z.read(name).decode("utf-8", "ignore")):
                xml = z.read("ppt/comments/" + target).decode("utf-8", "ignore")
                for t in re.findall(r"<a:t>(.*?)</a:t>", xml, re.S) + re.findall(r"<p:text>(.*?)</p:text>", xml, re.S):
                    out.append((order.get(m.group(1)), t.strip()))
    return sorted(out, key=lambda x: (x[0] or 0))


def red_runs(prs):
    """(slide, text) of runs coloured red -- the user's 'needs attention' marks."""
    out = []
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            frames = [sh.text_frame] if sh.has_text_frame else []
            if getattr(sh, "has_table", False) and sh.has_table:
                frames = [c.text_frame for r in sh.table.rows for c in r.cells]
            for tf in frames:
                for para in tf.paragraphs:
                    for run in para.runs:
                        try:
                            rgb = str(run.font.color.rgb) if run.font.color and run.font.color.type is not None else ""
                        except Exception:
                            rgb = ""
                        if len(rgb) == 6 and int(rgb[0:2], 16) >= 0xB0 and int(rgb[2:4], 16) <= 0x40 and int(rgb[4:6], 16) <= 0x40:
                            out.append((i, run.text))
    return out


def texts(slide):
    return [s.get("text") for s in slide["shapes"] if "text" in s]


def main(slug):
    folder = ROOT / "docs" / "decks" / slug
    base_path = folder / ".baseline" / (slug + ".generated.pptx")
    cur_path = folder / (slug + ".pptx")
    base, cur = describe(Presentation(base_path)), describe(Presentation(cur_path))
    print("# %s: edits since generation\n" % slug)
    if len(base) != len(cur):
        print("Slide count %d -> %d" % (len(base), len(cur)))
    base_titles = [s["title"] for s in base]
    for s in cur:
        match = next((b for b in base if b["title"] == s["title"]), None)
        if match is None:
            print("\n## Slide %d NEW or retitled: %r" % (s["n"], s["title"]))
            for sh in s["shapes"]:
                print("  +", sh)
            if s["notes"]:
                print("  notes:", s["notes"])
            continue
        lines = []
        if match["n"] != s["n"]:
            lines.append("moved from slide %d" % match["n"])
        bt, ct = texts(match), texts(s)
        for t in ct:
            if t not in bt:
                lines.append("text now: %r" % t)
        for t in bt:
            if t not in ct:
                lines.append("text was: %r" % t)
        btab = [sh.get("table") for sh in match["shapes"] if "table" in sh]
        ctab = [sh.get("table") for sh in s["shapes"] if "table" in sh]
        if btab != ctab:
            for ti, (a, b) in enumerate(zip(btab, ctab)):
                for ri in range(max(len(a), len(b))):
                    ra = a[ri] if ri < len(a) else None
                    rb = b[ri] if ri < len(b) else None
                    if ra != rb:
                        lines.append("table %d row %d: %r -> %r" % (ti + 1, ri, ra, rb))
            if len(btab) != len(ctab):
                lines.append("tables %d -> %d" % (len(btab), len(ctab)))
        bimg = [(sh["image"], sh["box"]) for sh in match["shapes"] if "image" in sh]
        cimg = [(sh["image"], sh["box"]) for sh in s["shapes"] if "image" in sh]
        if [i for i, _ in bimg] != [i for i, _ in cimg]:
            lines.append("images changed: %d -> %d (new/replaced: %s)" % (len(bimg), len(cimg), [i for i, _ in cimg if i not in [j for j, _ in bimg]]))
        elif [b for _, b in bimg] != [b for _, b in cimg]:
            lines.append("images moved/resized: %s -> %s" % ([b for _, b in bimg], [b for _, b in cimg]))
        # PowerPoint re-measures table heights on save; only a moved or resized non-table shape is a layout edit.
        bbox = sorted(sh["box"] for sh in match["shapes"] if "table" not in sh)
        cbox = sorted(sh["box"] for sh in s["shapes"] if "table" not in sh)
        if bbox != cbox and not any(l.startswith("images") for l in lines):
            lines.append("layout changed (%d -> %d shapes)" % (len(match["shapes"]), len(s["shapes"])))
        if s["notes"] != match["notes"]:
            lines.append("NOTES: %s" % s["notes"])
        if lines:
            print("\n## Slide %d: %s" % (s["n"], s["title"]))
            for l in lines:
                print("  -", l)
    cur_titles = [s["title"] for s in cur]
    for b in base:
        if b["title"] not in cur_titles:
            print("\n## Slide %d DELETED: %r" % (b["n"], b["title"]))
    for n, text in comments(cur_path):
        print("\nCOMMENT slide %s: %s" % (n, text))
    for n, text in red_runs(Presentation(cur_path)):
        print("RED TEXT slide %d: %s" % (n, text))


if __name__ == "__main__":
    main(sys.argv[1])
